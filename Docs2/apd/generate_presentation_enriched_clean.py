"""
Script propre de génération de la présentation enrichie.
Génère PRESENTATION_TP_DWWM_ENRICHED.pptx (version minimale fonctionnelle).
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor

ROUGE = RGBColor(172, 17, 21)
NOIR = RGBColor(23, 23, 23)
BLANC = RGBColor(255, 255, 255)
GRIS = RGBColor(237, 237, 237)

prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)

def add_blank():
    return prs.slides.add_slide(prs.slide_layouts[6])

def title_bar(slide, text):
    shape = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(10), Inches(1.1))
    shape.fill.solid(); shape.fill.fore_color.rgb = ROUGE
    shape.line.color.rgb = ROUGE
    tf = shape.text_frame; tf.text = text; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.font.size = Pt(30); p.font.bold = True; p.font.color.rgb = BLANC; p.alignment = PP_ALIGN.CENTER

def add_section_title(slide, text, y):
    box = slide.shapes.add_textbox(Inches(0.7), y, Inches(8.6), Inches(0.6))
    tf = box.text_frame; tf.text = text
    p = tf.paragraphs[0]; p.font.size = Pt(18); p.font.bold = True; p.font.color.rgb = ROUGE

def add_bullets(slide, items, x, y, w, h, size=16):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame; tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item; p.font.size = Pt(size); p.font.color.rgb = NOIR

def build():
    # Titre
    s = add_blank(); bg = s.background.fill; bg.solid(); bg.fore_color.rgb = ROUGE
    box = s.shapes.add_textbox(Inches(1), Inches(2.3), Inches(8), Inches(1.3))
    tf = box.text_frame; tf.text = 'DOSSIER DE PROJET - TP DWWM'
    p = tf.paragraphs[0]; p.font.size = Pt(48); p.font.bold = True; p.font.color.rgb = BLANC; p.alignment = PP_ALIGN.CENTER
    sub = s.shapes.add_textbox(Inches(1), Inches(3.9), Inches(8), Inches(0.9))
    ts = sub.text_frame; ts.text = 'APD • CoolBooking'
    ps = ts.paragraphs[0]; ps.font.size = Pt(24); ps.font.color.rgb = GRIS; ps.alignment = PP_ALIGN.CENTER

    # Sommaire
    s = add_blank(); title_bar(s, 'SOMMAIRE')
    add_bullets(s, [
        'Identité', 'Équipe', 'Architecture', 'Sécurité', 'Perspectives', 'Conclusion', 'Q&A'
    ], Inches(0.9), Inches(1.6), Inches(8.2), Inches(4.5), size=22)

    # Identité
    s = add_blank(); title_bar(s, 'IDENTITÉ')
    add_section_title(s, 'Objectif', Inches(1.4))
    add_bullets(s, [
        'Reconversion centrée produit & sécurité', 'MVP rapide puis durcissement', 'Cycle complet UX → API → DB'
    ], Inches(0.9), Inches(2.1), Inches(8.2), Inches(1.6))

    # Architecture
    s = add_blank(); title_bar(s, 'ARCHITECTURE')
    add_bullets(s, [
        'Next.js (SSR/RSC) + Strapi', 'Express + MariaDB + Cloudinary', 'Séparation frontend / backend', 'Scalabilité & testabilité'
    ], Inches(0.9), Inches(1.6), Inches(8.2), Inches(3.6))

    # Sécurité
    s = add_blank(); title_bar(s, 'SÉCURITÉ')
    add_section_title(s, 'Implémenté', Inches(1.4))
    add_bullets(s, ['Argon2id', 'JWT HttpOnly', 'Validation Joi', 'Requêtes paramétrées'], Inches(0.9), Inches(2.1), Inches(3.6), Inches(1.6))
    add_section_title(s, 'À renforcer', Inches(4.0))
    add_bullets(s, ['Rate limiting', 'Helmet', 'Logs sécurité', 'Refresh tokens'], Inches(0.9), Inches(4.6), Inches(3.6), Inches(1.6))

    # Perspectives
    s = add_blank(); title_bar(s, 'PERSPECTIVES')
    add_bullets(s, [
        'Tests Jest + CI/CD', 'Optimisation vidéo', 'Réservations Stripe', 'Internationalisation', 'PWA / Offline'
    ], Inches(0.9), Inches(1.6), Inches(8.2), Inches(3.6))

    # Conclusion
    s = add_blank(); title_bar(s, 'CONCLUSION')
    add_bullets(s, [
        'Base solide sécurité & structure', 'Vision claire évolutive', 'Livrables documentés', 'Focus valeur utilisateur'
    ], Inches(0.9), Inches(1.6), Inches(8.2), Inches(3.6))

    # Q&A
    s = add_blank(); title_bar(s, 'Q&A')
    box = s.shapes.add_textbox(Inches(1), Inches(2.8), Inches(8), Inches(2))
    tf = box.text_frame; tf.text = 'Merci !'
    p = tf.paragraphs[0]; p.font.size = Pt(60); p.font.bold = True; p.font.color.rgb = ROUGE; p.alignment = PP_ALIGN.CENTER

build()
prs.save('PRESENTATION_TP_DWWM_ENRICHED.pptx')
print('✅ Fichier généré : PRESENTATION_TP_DWWM_ENRICHED.pptx')