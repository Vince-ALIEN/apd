"""
Génération de la présentation PowerPoint pour le dossier TP DWWM
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor

# Couleurs de la charte APD
ROUGE_APD = RGBColor(172, 17, 21)  # #AC1115
NOIR = RGBColor(23, 23, 23)  # #171717
BLANC = RGBColor(255, 255, 255)
GRIS_CLAIR = RGBColor(237, 237, 237)

def create_presentation():
    """Crée la présentation PowerPoint complète"""
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)
    
    # Slide 1: Page de titre
    add_title_slide(prs)
    
    # Slide 2: Sommaire
    add_summary_slide(prs)
    
    # Slide 3: Contexte formation
    add_context_slide(prs)
    
    # Slide 4: Projet APD - Présentation
    add_apd_intro_slide(prs)
    
    # Slide 5: Projet APD - Technologies
    add_apd_tech_slide(prs)
    
    # Slide 6: Projet APD - Résultats
    add_apd_results_slide(prs)
    
    # Slide 7: Projet CoolBooking - Présentation
    add_coolbooking_intro_slide(prs)
    
    # Slide 8: Projet CoolBooking - Architecture
    add_coolbooking_arch_slide(prs)
    
    # Slide 9: Projet CoolBooking - Sécurité
    add_coolbooking_security_slide(prs)
    
    # Slide 10: Compétences CCP1
    add_ccp1_slide(prs)
    
    # Slide 11: Compétences CCP2
    add_ccp2_slide(prs)
    
    # Slide 12: Réalisations techniques
    add_technical_achievements_slide(prs)
    
    # Slide 13: Tests et qualité
    add_tests_slide(prs)
    
    # Slide 14: Sécurité OWASP
    add_security_slide(prs)
    
    # Slide 15: Perspectives
    add_perspectives_slide(prs)
    
    # Slide 16: Conclusion
    add_conclusion_slide(prs)
    
    # Sauvegarde
    prs.save('PRESENTATION_TP_DWWM.pptx')
    print("✅ Présentation PowerPoint générée : PRESENTATION_TP_DWWM.pptx")

def add_title_slide(prs):
    """Slide 1: Page de titre"""
    slide_layout = prs.slide_layouts[6]  # Blank
    slide = prs.slides.add_slide(slide_layout)
    
    # Fond rouge
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = ROUGE_APD
    
    # Titre principal
    title_box = slide.shapes.add_textbox(Inches(1), Inches(2.5), Inches(8), Inches(1))
    title_frame = title_box.text_frame
    title_frame.text = "DOSSIER DE PROJET"
    title_para = title_frame.paragraphs[0]
    title_para.font.size = Pt(54)
    title_para.font.bold = True
    title_para.font.color.rgb = BLANC
    title_para.alignment = PP_ALIGN.CENTER
    
    # Sous-titre
    subtitle_box = slide.shapes.add_textbox(Inches(1), Inches(3.7), Inches(8), Inches(0.8))
    subtitle_frame = subtitle_box.text_frame
    subtitle_frame.text = "Titre Professionnel Développeur Web et Web Mobile"
    subtitle_para = subtitle_frame.paragraphs[0]
    subtitle_para.font.size = Pt(28)
    subtitle_para.font.color.rgb = BLANC
    subtitle_para.alignment = PP_ALIGN.CENTER
    
    # Projets
    projects_box = slide.shapes.add_textbox(Inches(1), Inches(4.8), Inches(8), Inches(0.6))
    projects_frame = projects_box.text_frame
    projects_frame.text = "APD (Frontend Next.js) • CoolBooking (Backend Express)"
    projects_para = projects_frame.paragraphs[0]
    projects_para.font.size = Pt(20)
    projects_para.font.color.rgb = GRIS_CLAIR
    projects_para.alignment = PP_ALIGN.CENTER
    
    # Auteur et date
    author_box = slide.shapes.add_textbox(Inches(1), Inches(6.2), Inches(8), Inches(0.5))
    author_frame = author_box.text_frame
    author_frame.text = "Philippe Barbosa • Janvier 2025"
    author_para = author_frame.paragraphs[0]
    author_para.font.size = Pt(18)
    author_para.font.color.rgb = BLANC
    author_para.alignment = PP_ALIGN.CENTER

def add_summary_slide(prs):
    """Slide 2: Sommaire"""
    slide = add_slide_with_title(prs, "SOMMAIRE")
    
    content = [
        "1. Contexte de la formation",
        "2. Projet APD - Frontend Next.js",
        "3. Projet CoolBooking - Backend Express",
        "4. Compétences CCP1 & CCP2",
        "5. Réalisations techniques",
        "6. Tests et qualité",
        "7. Sécurité OWASP Top 10",
        "8. Perspectives et conclusion"
    ]
    
    add_bullet_list(slide, content, Inches(1.5), Inches(2), Inches(7), Inches(4.5))

def add_context_slide(prs):
    """Slide 3: Contexte formation"""
    slide = add_slide_with_title(prs, "CONTEXTE DE LA FORMATION")
    
    # Formation
    add_section_title(slide, "📚 Formation TP DWWM", Inches(1), Inches(2))
    formation_content = [
        "Titre Professionnel Développeur Web et Web Mobile (RNCP niveau 5)",
        "2 CCP : Frontend (CCP1) + Backend (CCP2)",
        "Durée : 12 semaines (8 semaines APD + 4 semaines CoolBooking)"
    ]
    add_bullet_list(slide, formation_content, Inches(1.2), Inches(2.5), Inches(7.5), Inches(1.5))
    
    # Objectifs
    add_section_title(slide, "🎯 Objectifs", Inches(1), Inches(4.2))
    objectifs_content = [
        "Démontrer la maîtrise complète du développement web",
        "Frontend : Maquettage, intégration responsive, animations",
        "Backend : API REST, sécurité, base de données"
    ]
    add_bullet_list(slide, objectifs_content, Inches(1.2), Inches(4.7), Inches(7.5), Inches(1.5))

def add_apd_intro_slide(prs):
    """Slide 4: Projet APD - Présentation"""
    slide = add_slide_with_title(prs, "PROJET APD - ASSOCIATION PATRIMOINE")
    
    # Contexte
    add_section_title(slide, "📍 Contexte", Inches(1), Inches(2))
    contexte = [
        "Client : Association Patrimoine de Doazit",
        "Mission : Préservation de l'église Saint-Jean Baptiste d'Aulès",
        "Problématique : Absence de présence digitale, difficulté à collecter des dons"
    ]
    add_bullet_list(slide, contexte, Inches(1.2), Inches(2.5), Inches(7.5), Inches(1.3))
    
    # Solution
    add_section_title(slide, "💡 Solution", Inches(1), Inches(4))
    solution = [
        "Site web immersif Next.js 15 avec animations GSAP",
        "Vidéo background émotionnelle + galerie photos",
        "Blog éditorial + interviews vidéo + section partenaires",
        "CTA 'Faire un don' permanent et optimisé pour conversions"
    ]
    add_bullet_list(slide, solution, Inches(1.2), Inches(4.5), Inches(7.5), Inches(1.8))

def add_apd_tech_slide(prs):
    """Slide 5: Projet APD - Technologies"""
    slide = add_slide_with_title(prs, "PROJET APD - STACK TECHNIQUE")
    
    # Colonne gauche : Frontend
    add_section_title(slide, "⚛️ Frontend", Inches(1), Inches(2))
    frontend = [
        "Next.js 15.5.3 (App Router)",
        "React 19.1.0",
        "Tailwind CSS 4",
        "GSAP 3.13 (animations)",
        "16 composants modulaires"
    ]
    add_bullet_list(slide, frontend, Inches(1.2), Inches(2.5), Inches(3.5), Inches(2.5))
    
    # Colonne droite : Backend & Déploiement
    add_section_title(slide, "🗄️ Backend CMS", Inches(5.5), Inches(2))
    backend = [
        "Strapi v5 (headless CMS)",
        "PostgreSQL 14+",
        "Cloudinary (CDN médias)",
        "3 Single Types + 3 Collections"
    ]
    add_bullet_list(slide, backend, Inches(5.7), Inches(2.5), Inches(3.5), Inches(2))
    
    add_section_title(slide, "🚀 Déploiement", Inches(5.5), Inches(4.7))
    deploy = [
        "Frontend : Vercel",
        "Backend : Railway",
        "Médias : Cloudinary CDN"
    ]
    add_bullet_list(slide, deploy, Inches(5.7), Inches(5.2), Inches(3.5), Inches(1.3))

def add_apd_results_slide(prs):
    """Slide 6: Projet APD - Résultats"""
    slide = add_slide_with_title(prs, "PROJET APD - RÉSULTATS")
    
    # Lighthouse scores
    add_section_title(slide, "📊 Lighthouse Scores (Desktop)", Inches(1), Inches(2))
    scores_box = slide.shapes.add_textbox(Inches(1.2), Inches(2.5), Inches(7.5), Inches(1.5))
    scores_frame = scores_box.text_frame
    scores_frame.word_wrap = True
    
    scores_text = "Performance: 92/100  •  Accessibility: 90/100  •  Best Practices: 100/100  •  SEO: 100/100"
    p = scores_frame.paragraphs[0]
    p.text = scores_text
    p.font.size = Pt(18)
    p.font.bold = True
    p.alignment = PP_ALIGN.CENTER
    
    # Fonctionnalités clés
    add_section_title(slide, "✨ Fonctionnalités clés", Inches(1), Inches(4.2))
    features = [
        "Navigation fluide entre sections avec GSAP ScrollTrigger",
        "Galerie responsive (9 images) avec lazy loading",
        "Blog dynamique avec pages dynamiques [slug]",
        "Formulaire contact partenaires avec API email",
        "Optimisations : Meta tags, sitemap, robots.txt"
    ]
    add_bullet_list(slide, features, Inches(1.2), Inches(4.7), Inches(7.5), Inches(2.3))

def add_coolbooking_intro_slide(prs):
    """Slide 7: Projet CoolBooking - Présentation"""
    slide = add_slide_with_title(prs, "PROJET COOLBOOKING - API REST")
    
    # Contexte
    add_section_title(slide, "📍 Contexte", Inches(1), Inches(2))
    contexte = [
        "Problématique : Airbnb prélève 15-20% de commission",
        "Besoin : Alternative pour propriétaires indépendants",
        "Cible : Petites agences locales, propriétaires multi-biens"
    ]
    add_bullet_list(slide, contexte, Inches(1.2), Inches(2.5), Inches(7.5), Inches(1.3))
    
    # Solution
    add_section_title(slide, "💡 Solution technique", Inches(1), Inches(4))
    solution = [
        "API REST Express 5 avec architecture modulaire",
        "Authentification JWT + hash Argon2id",
        "Upload images Cloudinary (5 par annonce)",
        "Base MariaDB avec pool de connexions optimisé",
        "Validation Joi + gestion erreurs globale"
    ]
    add_bullet_list(slide, solution, Inches(1.2), Inches(4.5), Inches(7.5), Inches(2))

def add_coolbooking_arch_slide(prs):
    """Slide 8: Projet CoolBooking - Architecture"""
    slide = add_slide_with_title(prs, "COOLBOOKING - ARCHITECTURE")
    
    # Pattern MVC
    add_section_title(slide, "🏗️ Architecture en couches (Repository Pattern)", Inches(1), Inches(2))
    
    # Diagramme ASCII
    arch_box = slide.shapes.add_textbox(Inches(2), Inches(2.7), Inches(6), Inches(2.5))
    arch_frame = arch_box.text_frame
    arch_frame.word_wrap = True
    
    arch_text = """Routes (Express Router)
        ↓
Controllers (gestion requêtes/réponses)
        ↓
Services (logique métier, hash Argon2)
        ↓
Repositories (requêtes SQL paramétrées)
        ↓
Database (Pool mysql2 - 10 connexions)"""
    
    p = arch_frame.paragraphs[0]
    p.text = arch_text
    p.font.size = Pt(16)
    p.font.name = "Courier New"
    p.alignment = PP_ALIGN.CENTER
    
    # Endpoints
    add_section_title(slide, "📡 Endpoints", Inches(1), Inches(5.5))
    endpoints = [
        "/users : CRUD + auth (register, login, logout, dashboard)",
        "/rentals : CRUD annonces avec upload 5 images max"
    ]
    add_bullet_list(slide, endpoints, Inches(1.2), Inches(6), Inches(7.5), Inches(1))

def add_coolbooking_security_slide(prs):
    """Slide 9: Projet CoolBooking - Sécurité"""
    slide = add_slide_with_title(prs, "COOLBOOKING - SÉCURITÉ")
    
    # Colonne gauche
    add_section_title(slide, "🔒 Authentification", Inches(1), Inches(2))
    auth = [
        "Argon2id : hash passwords (résistant GPU)",
        "JWT : tokens signés (expiration 24h)",
        "HttpOnly cookies : protection XSS",
        "SameSite: Strict : protection CSRF"
    ]
    add_bullet_list(slide, auth, Inches(1.2), Inches(2.5), Inches(3.8), Inches(2))
    
    # Colonne droite
    add_section_title(slide, "🛡️ Protection données", Inches(5), Inches(2))
    protection = [
        "Requêtes paramétrées (anti SQL injection)",
        "Validation Joi (email, password, phone)",
        "CORS configuré",
        "Gestion erreurs globale"
    ]
    add_bullet_list(slide, protection, Inches(5.2), Inches(2.5), Inches(3.8), Inches(2))
    
    # Bas
    add_section_title(slide, "✅ Bonnes pratiques", Inches(1), Inches(5))
    bp = [
        "Séparation des responsabilités (Controller/Service/Repository)",
        "Variables env (.env non committé)",
        "Pool connexions DB optimisé (release auto)",
        "Suppression fichiers temporaires après upload Cloudinary"
    ]
    add_bullet_list(slide, bp, Inches(1.2), Inches(5.5), Inches(7.5), Inches(1.5))

def add_ccp1_slide(prs):
    """Slide 10: Compétences CCP1"""
    slide = add_slide_with_title(prs, "COMPÉTENCES CCP1 - FRONTEND")
    
    # CCP1-A
    add_section_title(slide, "🎨 CCP1-A : Maquetter une application", Inches(1), Inches(2))
    ccp1a = [
        "Wireframes desktop/mobile (diagrammes ASCII)",
        "Charte graphique (couleurs, typographies, composants)",
        "User stories (23 APD) + personas détaillés (4)"
    ]
    add_bullet_list(slide, ccp1a, Inches(1.2), Inches(2.5), Inches(7.5), Inches(1.2))
    
    # CCP1-B
    add_section_title(slide, "📱 CCP1-B : Interface statique et adaptable", Inches(1), Inches(3.9))
    ccp1b = [
        "Intégration responsive Tailwind CSS (mobile-first)",
        "Accessibilité (contraste AAA, navigation clavier, alt images)",
        "SEO optimisé (meta tags, sitemap, robots.txt)"
    ]
    add_bullet_list(slide, ccp1b, Inches(1.2), Inches(4.4), Inches(7.5), Inches(1.2))
    
    # CCP1-C
    add_section_title(slide, "⚡ CCP1-C : Interface dynamique", Inches(1), Inches(5.8))
    ccp1c = [
        "Animations GSAP (ScrollTrigger, timelines, reveal effects)",
        "Hooks React (useSiteData, useCurrentSection) + Context API",
        "Fetch API Strapi (single types, collections) + Next.js dynamic routes"
    ]
    add_bullet_list(slide, ccp1c, Inches(1.2), Inches(6.3), Inches(7.5), Inches(1.2))

def add_ccp2_slide(prs):
    """Slide 11: Compétences CCP2"""
    slide = add_slide_with_title(prs, "COMPÉTENCES CCP2 - BACKEND")
    
    # CCP2-A
    add_section_title(slide, "🗄️ CCP2-A : Créer une base de données", Inches(1), Inches(2))
    ccp2a = [
        "Modélisation MCD/MLD (users, rentals)",
        "Dictionnaire de données (types, contraintes)",
        "Migrations Strapi PostgreSQL + tables MariaDB"
    ]
    add_bullet_list(slide, ccp2a, Inches(1.2), Inches(2.5), Inches(7.5), Inches(1.2))
    
    # CCP2-B
    add_section_title(slide, "🔌 CCP2-B : Composants d'accès aux données", Inches(1), Inches(3.9))
    ccp2b = [
        "Repository pattern (encapsulation SQL)",
        "Requêtes paramétrées mysql2 (protection injection)",
        "Pool connexions optimisé (10 max) + ORM Strapi"
    ]
    add_bullet_list(slide, ccp2b, Inches(1.2), Inches(4.4), Inches(7.5), Inches(1.2))
    
    # CCP2-C
    add_section_title(slide, "🚀 CCP2-C : Développer le backend", Inches(1), Inches(5.8))
    ccp2c = [
        "API REST Express (CRUD users, rentals) + architecture modulaire",
        "Authentification JWT + hash Argon2id + validation Joi",
        "Upload Cloudinary + CORS + gestion erreurs globale"
    ]
    add_bullet_list(slide, ccp2c, Inches(1.2), Inches(6.3), Inches(7.5), Inches(1.2))

def add_technical_achievements_slide(prs):
    """Slide 12: Réalisations techniques"""
    slide = add_slide_with_title(prs, "RÉALISATIONS TECHNIQUES")
    
    # APD
    add_section_title(slide, "🎯 APD - Animations GSAP", Inches(1), Inches(2))
    apd_tech = [
        "ScrollTrigger pour animations au scroll (reveal boxes)",
        "Timelines pour séquences complexes (cascade titre → texte)",
        "Pulse animations (bouton donation, icône cœur)"
    ]
    add_bullet_list(slide, apd_tech, Inches(1.2), Inches(2.5), Inches(7.5), Inches(1.3))
    
    # CoolBooking
    add_section_title(slide, "🔐 CoolBooking - Authentification JWT", Inches(1), Inches(4))
    cb_tech = [
        "verifyPassword : Argon2 verify + génération JWT + cookie HttpOnly",
        "verifyToken : middleware protection routes privées (/dashboard)",
        "Repository pattern : séparation logique métier / accès données",
        "Upload Cloudinary : 5 images max + suppression fichiers locaux"
    ]
    add_bullet_list(slide, cb_tech, Inches(1.2), Inches(4.5), Inches(7.5), Inches(2))

def add_tests_slide(prs):
    """Slide 13: Tests et qualité"""
    slide = add_slide_with_title(prs, "TESTS ET QUALITÉ")
    
    # APD Tests
    add_section_title(slide, "✅ APD - Tests manuels", Inches(1), Inches(2))
    apd_tests = [
        "10 scénarios navigation (scroll sections, animations, CTA)",
        "Tests responsive (iPhone SE, iPad, Desktop 1920x1080, 4K)",
        "Lighthouse Desktop : Perf 92, A11y 90, BP 100, SEO 100"
    ]
    add_bullet_list(slide, apd_tests, Inches(1.2), Inches(2.5), Inches(7.5), Inches(1.3))
    
    # CoolBooking Tests
    add_section_title(slide, "✅ CoolBooking - Collection Postman", Inches(1), Inches(4))
    cb_tests = [
        "18 requêtes testées (CRUD users + rentals + auth)",
        "Tests succès (200, 201) + erreurs (400, 401, 404, 409)",
        "Temps réponse : GET < 60ms, POST register 450ms, POST rentals 1200ms",
        "À implémenter : Tests unitaires Jest + CI/CD GitHub Actions"
    ]
    add_bullet_list(slide, cb_tests, Inches(1.2), Inches(4.5), Inches(7.5), Inches(2))

def add_security_slide(prs):
    """Slide 14: Sécurité OWASP"""
    slide = add_slide_with_title(prs, "SÉCURITÉ - OWASP TOP 10 (2021)")
    
    content = [
        "A01 - Broken Access Control : ✅ Middleware verifyToken, rôles Strapi",
        "A02 - Cryptographic Failures : ✅ Argon2id, JWT secrets .env, HTTPS",
        "A03 - Injection : ✅ Requêtes paramétrées, validation Joi",
        "A04 - Insecure Design : ✅ Architecture en couches, séparation frontend/backend",
        "A05 - Security Misconfiguration : ✅ CORS, SameSite cookies, headers Next.js",
        "A06 - Vulnerable Components : ✅ npm audit, Dependabot, versions récentes",
        "A07 - Auth Failures : ✅ JWT expiration 24h, Argon2 anti-bruteforce",
        "A09 - Logging Failures : ⚠️ À implémenter : Winston logger, Sentry"
    ]
    add_bullet_list(slide, content, Inches(1), Inches(2), Inches(8), Inches(4.5), font_size=14)

def add_perspectives_slide(prs):
    """Slide 15: Perspectives"""
    slide = add_slide_with_title(prs, "PERSPECTIVES D'ÉVOLUTION")
    
    # APD
    add_section_title(slide, "📈 APD - Court terme (3 mois)", Inches(1), Inches(2))
    apd_persp = [
        "Compression vidéo (HandBrake RF 32) → -40% poids",
        "Newsletter (Strapi + Nodemailer) + calendrier événements",
        "Multilingue (FR/EN/ES) + mode sombre",
        "Analytics (Plausible RGPD-friendly) + heatmaps Hotjar"
    ]
    add_bullet_list(slide, apd_persp, Inches(1.2), Inches(2.5), Inches(7.5), Inches(1.8))
    
    # CoolBooking
    add_section_title(slide, "📈 CoolBooking - Court terme (3 mois)", Inches(1), Inches(4.5))
    cb_persp = [
        "Système réservations (calendrier, paiements Stripe)",
        "Filtres annonces (prix, localisation, couchages) + messagerie",
        "Tests Jest + CI/CD + documentation Swagger",
        "Rate limiting + Helmet.js + Winston logger + 2FA"
    ]
    add_bullet_list(slide, cb_persp, Inches(1.2), Inches(5), Inches(7.5), Inches(1.8))

def add_conclusion_slide(prs):
    """Slide 16: Conclusion"""
    slide = add_slide_with_title(prs, "CONCLUSION")
    
    # Compétences
    add_section_title(slide, "💪 Compétences démontrées", Inches(1), Inches(2))
    competences = [
        "Frontend : Next.js 15, React 19, GSAP, Tailwind, responsive, accessibilité, SEO",
        "Backend : Express 5, MariaDB, JWT, Argon2, Repository pattern, API REST",
        "Sécurité : OWASP Top 10, requêtes paramétrées, validation, HTTPS, cookies HttpOnly",
        "Qualité : Tests Lighthouse 92/100, collection Postman 18 requêtes, documentation"
    ]
    add_bullet_list(slide, competences, Inches(1.2), Inches(2.5), Inches(7.5), Inches(2), font_size=15)
    
    # Valeur ajoutée
    add_section_title(slide, "🎯 Valeur ajoutée", Inches(1), Inches(4.8))
    valeur = [
        "Projet complet : du maquettage à la production (Vercel, Railway)",
        "Architecture professionnelle : modulaire, testable, évolutive",
        "Veille technologique : Next.js 15, React 19, Express 5, Argon2id",
        "Prêt pour environnement professionnel"
    ]
    add_bullet_list(slide, valeur, Inches(1.2), Inches(5.3), Inches(7.5), Inches(1.5), font_size=15)

# Fonctions utilitaires

def add_slide_with_title(prs, title_text):
    """Ajoute une slide avec titre standard"""
    slide_layout = prs.slide_layouts[6]  # Blank
    slide = prs.slides.add_slide(slide_layout)
    
    # Barre titre rouge
    title_shape = slide.shapes.add_shape(
        1,  # Rectangle
        Inches(0), Inches(0),
        Inches(10), Inches(1.2)
    )
    title_shape.fill.solid()
    title_shape.fill.fore_color.rgb = ROUGE_APD
    title_shape.line.color.rgb = ROUGE_APD
    
    # Texte titre
    title_frame = title_shape.text_frame
    title_frame.text = title_text
    title_frame.margin_top = Inches(0.2)
    title_frame.margin_bottom = Inches(0.2)
    title_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    
    p = title_frame.paragraphs[0]
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = BLANC
    p.alignment = PP_ALIGN.CENTER
    
    return slide

def add_section_title(slide, text, left, top):
    """Ajoute un sous-titre de section"""
    box = slide.shapes.add_textbox(left, top, Inches(7.5), Inches(0.4))
    frame = box.text_frame
    frame.text = text
    p = frame.paragraphs[0]
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ROUGE_APD

def add_bullet_list(slide, items, left, top, width, height, font_size=16):
    """Ajoute une liste à puces"""
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    
    for i, item in enumerate(items):
        if i == 0:
            p = frame.paragraphs[0]
        else:
            p = frame.add_paragraph()
        
        p.text = item
        p.font.size = Pt(font_size)
        p.font.color.rgb = NOIR
        p.level = 0
        p.space_before = Pt(6)

if __name__ == "__main__":
    create_presentation()
